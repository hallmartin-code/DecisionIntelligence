"""Renderer tests: the .docx must match `templates/report_structure.md`."""

from __future__ import annotations

from datetime import date

import pytest
from docx import Document
from docx.oxml.ns import qn

from pitch_analyzer.models import CATEGORIES, CATEGORY_KEYS, AnalysisResult
from pitch_analyzer.render import (
    CRIMSON,
    HEADER_FILL,
    NAVY,
    PANEL_FILL,
    render_report,
)

STAMP = date(2026, 8, 31)
INCH = 914400


@pytest.fixture
def report_path(tmp_path, analysis_result):
    return render_report(
        analysis_result, tmp_path / "report.docx", generated_on=STAMP
    )


@pytest.fixture
def report(report_path):
    return Document(str(report_path))


def _text(document) -> str:
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _fill(cell) -> str | None:
    properties = cell._tc.tcPr
    if properties is None:
        return None
    shading = properties.find(qn("w:shd"))
    return shading.get(qn("w:fill")) if shading is not None else None


def _headings(document, size_pt: float) -> list[str]:
    return [
        p.text
        for p in document.paragraphs
        if p.runs
        and p.runs[0].font.size
        and abs(p.runs[0].font.size.pt - size_pt) < 0.01
        and p.runs[0].font.bold
    ]


# --------------------------------------------------------------------------- #
# Output basics
# --------------------------------------------------------------------------- #


def test_render_writes_a_readable_docx(tmp_path, analysis_result):
    output = tmp_path / "out.docx"

    returned = render_report(analysis_result, output, generated_on=STAMP)

    assert returned == output
    assert output.exists()
    assert output.stat().st_size > 10 * 1024
    assert output.read_bytes()[:2] == b"PK"  # a zip container, i.e. real OOXML


def test_page_setup_matches_the_specification(report):
    section = report.sections[0]

    assert round(section.page_width / INCH, 2) == 8.5
    assert round(section.page_height / INCH, 2) == 11.0
    assert round(section.left_margin / INCH, 2) == 0.88
    assert round(section.top_margin / INCH, 2) == 0.83


def test_footer_carries_the_company_and_a_page_field(report):
    footer = report.sections[0].footer.paragraphs[0]

    assert "Decision Intelligence Assessment" in footer.text
    assert "Acme Robotics" in footer.text
    # This goes to an investment committee; the marker stays even though the
    # house footer pattern does not call for one.
    assert "Confidential" in footer.text
    assert "by TEN Capital Network" in footer.text
    instructions = footer._p.findall(".//" + qn("w:instrText"))
    assert [i.text.strip() for i in instructions] == ["PAGE", "NUMPAGES"]


def test_the_footer_follows_the_house_standard(report):
    """Open Sans 7pt, centred - so a TEN document is recognisable as one."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    footer = report.sections[0].footer.paragraphs[0]

    assert footer.alignment == WD_ALIGN_PARAGRAPH.CENTER
    fonts = {run.font.name for run in footer.runs if run.font.name}
    sizes = {run.font.size for run in footer.runs if run.font.size}
    assert fonts == {"Open Sans"}, fonts
    assert sizes == {Pt(7)}, sizes


def test_the_footer_carries_the_logo(report_path):
    """The mark has to be embedded in the file, not merely referenced."""
    import re
    import zipfile

    from pitch_analyzer.render import LOGO_HEIGHT, LOGO_PATH, LOGO_WIDTH

    with zipfile.ZipFile(report_path) as archive:
        media = [n for n in archive.namelist() if n.startswith("word/media/")]
        assert media, "no image was embedded in the document"
        assert any(
            archive.read(name) == LOGO_PATH.read_bytes() for name in media
        ), "the embedded image is not the TEN Capital mark"

        footer_xml = archive.read("word/footer1.xml").decode("utf-8")

    assert "a:blip" in footer_xml, "the mark is not placed in the footer"
    extent = re.findall(r'wp:extent cx="(\d+)" cy="(\d+)"', footer_xml)
    assert (str(LOGO_WIDTH), str(LOGO_HEIGHT)) in extent, extent


def test_a_missing_logo_does_not_fail_the_report(tmp_path, analysis_result, monkeypatch):
    """The analysis is the deliverable; the mark is decoration."""
    import docx

    from pitch_analyzer import render

    monkeypatch.setattr(render, "LOGO_PATH", tmp_path / "absent.png")
    destination = tmp_path / "no_logo.docx"

    render.render_report(analysis_result, destination)

    footer = docx.Document(destination).sections[0].footer.paragraphs[0]
    assert "Decision Intelligence Assessment" in footer.text


# --------------------------------------------------------------------------- #
# Document skeleton
# --------------------------------------------------------------------------- #


def test_all_seven_parts_are_present_in_order(report):
    assert _headings(report, 15.0) == [
        "Executive Summary",
        "Decision Intelligence Assessment",
        "Decision Scenario Analysis",
        "Investment Committee View",
        "Decision Intelligence Scorecard",
        "Final Recommendation",
        "Summary Investment Memo",
    ]


def test_all_ten_categories_appear_numbered_and_scored(report):
    h2 = _headings(report, 12.0)

    for number, (_key, title) in enumerate(CATEGORIES, start=1):
        assert any(
            heading.startswith(f"{number}.  {title} — Score ") for heading in h2
        ), f"missing assessment section: {title}"


def test_masthead_and_metadata(report):
    text = _text(report)

    assert "TEN CAPITAL GROUP" in text
    assert "INVESTMENT COMMITTEE" in text
    assert "Acme Robotics" in text
    assert "Warehouse picking automation for mid-size facilities" in text
    assert "31 August 2026" in text
    assert "Not disclosed" in text  # absent fields are stated, not blank


def test_scoring_fairness_note_is_carried(report):
    assert "A note on scoring this document fairly" in _text(report)


# --------------------------------------------------------------------------- #
# Tables
# --------------------------------------------------------------------------- #


def test_grid_tables_use_the_navy_header(report):
    grids = [t for t in report.tables if len(t.columns) > 1]

    assert len(grids) >= 10
    for table in grids:
        assert _fill(table.rows[0].cells[0]) == HEADER_FILL


def test_callouts_are_shaded_panels(report):
    callouts = [t for t in report.tables if len(t.columns) == 1]

    assert len(callouts) >= 4
    for table in callouts:
        assert _fill(table.rows[0].cells[0]) == PANEL_FILL
    titles = [t.cell(0, 0).paragraphs[0].text for t in callouts]
    assert any(title.startswith("RECOMMENDATION:") for title in titles)
    assert any("BASIS OF THIS ANALYSIS" in title for title in titles)


def test_scorecard_arithmetic_closes(report, analysis_result):
    scorecard = next(
        t for t in report.tables
        if t.rows[0].cells[0].text == "Category" and len(t.columns) == 5
    )

    assert len(scorecard.rows) == 12  # header + 10 categories + total
    weighted_sum = 0.0
    for row in scorecard.rows[1:-1]:
        score = int(row.cells[1].text.split("/")[0].strip())
        weight = int(row.cells[2].text.rstrip("%"))
        weighted = float(row.cells[3].text)
        assert weighted == pytest.approx(round(score * weight / 100, 2))
        weighted_sum += weighted

    total_row = scorecard.rows[-1]
    assert total_row.cells[0].text == "WEIGHTED OVERALL SCORE"
    assert total_row.cells[2].text == "100%"
    assert float(total_row.cells[3].text) == pytest.approx(weighted_sum, abs=0.01)
    assert weighted_sum == pytest.approx(analysis_result.weighted_overall, abs=0.01)


def test_risk_register_is_numbered_and_prefixed(report):
    register = next(
        t for t in report.tables
        if t.rows[0].cells[-1].text == "Mitigation strategy"
    )

    assert [row.cells[0].text for row in register.rows[1:]] == ["1", "2", "3"]
    assert register.rows[1].cells[1].text.startswith("REGULATORY — ")


def test_scenario_table_accumulates(report):
    scenarios = next(
        t for t in report.tables if t.rows[0].cells[0].text == "Scenario"
    )

    cumulative = [row.cells[4].text for row in scenarios.rows[1:]]
    assert cumulative == ["2.20x", "2.83x", "2.92x"]


def test_memo_lists_all_nine_fields(report):
    memo = report.tables[-1]

    assert memo.rows[0].cells[0].text == "Field"
    assert len(memo.rows) == 10
    assert memo.rows[1].cells[0].text == "Company"
    assert memo.rows[-1].cells[0].text == "Decision"


def test_comparative_context_is_omitted_when_empty(report):
    assert "Comparative context" not in _text(report)


def test_comparative_context_renders_when_supplied(tmp_path, analysis_payload):
    analysis_payload["comparative_context"] = [
        {
            "company": "Peer Co",
            "di_score": "6.5 / 10",
            "confidence": "74%",
            "central_issue": "One answerable question.",
        }
    ]
    analysis = AnalysisResult.model_validate(analysis_payload)

    path = render_report(analysis, tmp_path / "peers.docx", generated_on=STAMP)
    text = _text(Document(str(path)))

    assert "Comparative context across this cycle" in text
    assert "Peer Co" in text


# --------------------------------------------------------------------------- #
# Colour conventions
# --------------------------------------------------------------------------- #


def test_adverse_bullets_are_set_in_crimson(report):
    crimson = [
        run.text
        for p in report.paragraphs
        for run in p.runs
        if run.font.color and run.font.color.rgb == CRIMSON
    ]

    assert any("unsourced" in text for text in crimson)
    assert any("INVESTIGATE FURTHER" in text for text in crimson)


def test_part_headings_are_navy(report):
    navy = [
        p.text
        for p in report.paragraphs
        if p.runs and p.runs[0].font.color and p.runs[0].font.color.rgb == NAVY
    ]

    assert "Executive Summary" in navy
    assert "Acme Robotics" in navy


# --------------------------------------------------------------------------- #
# Degenerate input
# --------------------------------------------------------------------------- #


def test_renders_with_empty_optional_tables(tmp_path, analysis_payload):
    for key in ("market_sizing", "competitors", "evidence_quality",
                "sensitivity", "risk_register", "assumptions"):
        analysis_payload[key] = []
    analysis_payload["basis_of_analysis"] = None
    analysis_payload["expected_outcome_callout"] = None
    analysis = AnalysisResult.model_validate(analysis_payload)

    path = render_report(analysis, tmp_path / "sparse.docx", generated_on=STAMP)

    assert path.stat().st_size > 5 * 1024
    assert len(_headings(Document(str(path)), 15.0)) == 7


def test_markup_characters_survive(tmp_path, analysis_payload):
    analysis_payload["verdict_paragraph"] = "Margins >50% & <2yr payback for R&D"
    analysis = AnalysisResult.model_validate(analysis_payload)

    path = render_report(analysis, tmp_path / "escaped.docx", generated_on=STAMP)

    assert "Margins >50% & <2yr payback" in _text(Document(str(path)))
