"""Ingestion tests: a generated 5-page PDF and a generated 5-slide PPTX."""

from __future__ import annotations

import pytest

from pitch_analyzer.ingest import (
    DeckContent,
    UnsupportedDeckError,
    _sample_indices,
    guess_company_name,
    ingest,
)

SLIDE_TITLES = [
    "Acme Robotics",
    "The Problem",
    "Our Solution",
    "Traction",
    "The Ask",
]


@pytest.fixture(scope="module")
def sample_pdf(tmp_path_factory) -> str:
    """A tiny 5-page PDF built with ReportLab."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path = tmp_path_factory.mktemp("decks") / "sample.pdf"
    pdf = canvas.Canvas(str(path), pagesize=letter)
    for index, title in enumerate(SLIDE_TITLES, start=1):
        pdf.setFont("Helvetica-Bold", 24)
        pdf.drawString(72, 700, title)
        pdf.setFont("Helvetica", 12)
        pdf.drawString(72, 660, f"Body copy for page {index} of the deck.")
        pdf.showPage()
    pdf.save()
    return str(path)


@pytest.fixture(scope="module")
def sample_pptx(tmp_path_factory) -> str:
    """A tiny 5-slide PPTX built with python-pptx."""
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path_factory.mktemp("decks") / "sample.pptx"
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    for index, title in enumerate(SLIDE_TITLES, start=1):
        slide = presentation.slides.add_slide(blank)
        box = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(8), Inches(2)
        )
        frame = box.text_frame
        frame.text = title
        frame.add_paragraph().text = f"Body copy for slide {index} of the deck."
    presentation.save(str(path))
    return str(path)


def test_pdf_slide_count_and_markers(sample_pdf):
    content = ingest(sample_pdf, include_images=False)

    assert isinstance(content, DeckContent)
    assert content.slide_count == 5
    for number in range(1, 6):
        assert f"[Slide {number}]" in content.text
    assert "Acme Robotics" in content.text
    assert "Body copy for page 3" in content.text
    assert content.images == []


def test_pptx_slide_count_and_markers(sample_pptx):
    content = ingest(sample_pptx, include_images=False)

    assert content.slide_count == 5
    for number in range(1, 6):
        assert f"[Slide {number}]" in content.text
    assert "The Ask" in content.text
    assert content.images == []


def test_image_extraction_is_capped(sample_pdf):
    content = ingest(sample_pdf, include_images=True)

    # Rendering depends on an optional native backend; when it is available we
    # must respect the cap, when it is not we must degrade to text-only.
    assert len(content.images) <= 5
    assert all(isinstance(blob, bytes) and blob for blob in content.images)


def test_unsupported_extension_rejected(tmp_path):
    bad = tmp_path / "deck.key"
    bad.write_text("not a deck")

    with pytest.raises(UnsupportedDeckError):
        ingest(bad)


def test_empty_deck_raises_runtime_error(tmp_path):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path = tmp_path / "blank.pdf"
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.showPage()
    pdf.save()

    with pytest.raises(RuntimeError, match="No text could be extracted."):
        ingest(path, include_images=False)


@pytest.mark.parametrize(
    ("total", "limit", "expected"),
    [
        (3, 10, 3),
        (10, 10, 10),
        (40, 10, 10),
        (0, 10, 0),
    ],
)
def test_sample_indices_respects_cap(total, limit, expected):
    picked = _sample_indices(total, limit)

    assert len(picked) == expected
    if total:
        assert 0 in picked
        assert total - 1 in picked
    assert all(0 <= index < total for index in picked)


def test_guess_company_name_uses_title_slide(sample_pdf):
    content = ingest(sample_pdf, include_images=False)

    assert guess_company_name(content.text) == "Acme Robotics"


def test_guess_company_name_skips_boilerplate():
    text = "[Slide 1]\nCONFIDENTIAL\n2026\nNorthwind Bio\n\n[Slide 2]\nProblem"

    assert guess_company_name(text) == "Northwind Bio"
