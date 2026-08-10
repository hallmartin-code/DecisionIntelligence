"""PDF / PPTX text + image extraction.

Produces a `DeckContent` that the analysis layer feeds to the model. Slide text
is concatenated into one string with `[Slide N]` markers so the model can cite
slide numbers; images are JPEG-encoded and capped so vision requests stay
within sane token budgets.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, List

MAX_IMAGES = 10
MAX_IMAGE_WIDTH = 1024
PDF_RENDER_DPI = 100


@dataclass
class DeckContent:
    """Everything extracted from a deck."""

    text: str
    slide_count: int
    images: List[bytes] = field(default_factory=list)


class UnsupportedDeckError(ValueError):
    """Raised for a file extension the ingester cannot read."""


def ingest(path: str | Path, include_images: bool = True) -> DeckContent:
    """Extract text (and optionally images) from a `.pdf` or `.pptx` deck."""
    deck_path = Path(path)
    suffix = deck_path.suffix.lower()
    if suffix == ".pdf":
        content = extract_pdf(deck_path, include_images=include_images)
    elif suffix == ".pptx":
        content = extract_pptx(deck_path, include_images=include_images)
    else:
        raise UnsupportedDeckError(
            f"Unsupported file type '{suffix or deck_path.name}'. Expected .pdf or .pptx."
        )

    if not _has_body_text(content.text):
        raise RuntimeError("No text could be extracted.")
    return content


_MARKER_LINE = re.compile(r"^\[(Slide|Table) \d+\]$")


def _has_body_text(text: str) -> bool:
    """True if the deck yielded anything beyond the structural markers."""
    return any(
        line.strip() and not _MARKER_LINE.match(line.strip())
        for line in text.splitlines()
    )


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #


def extract_pdf(path: str | Path, include_images: bool = True) -> DeckContent:
    """Extract per-page text, tables, and page screenshots from a PDF."""
    import pdfplumber

    chunks: list[str] = []
    page_images: list[bytes] = []

    with pdfplumber.open(str(path)) as pdf:
        pages = pdf.pages
        slide_count = len(pages)
        wanted = _sample_indices(slide_count, MAX_IMAGES) if include_images else set()

        for index, page in enumerate(pages):
            body = page.extract_text() or ""
            table_text = _format_tables(page.extract_tables() or [])
            chunks.append(_slide_block(index + 1, body, table_text))

            if index in wanted:
                jpeg = _render_pdf_page(page)
                if jpeg is not None:
                    page_images.append(jpeg)

    return DeckContent(
        text="\n\n".join(chunks), slide_count=slide_count, images=page_images
    )


def _render_pdf_page(page) -> bytes | None:
    """Rasterise one pdfplumber page to JPEG bytes, or None if unavailable."""
    try:
        rendered = page.to_image(resolution=PDF_RENDER_DPI)
        return _to_jpeg(rendered.original)
    except Exception:
        # Rendering needs an optional native backend; text-only analysis is a
        # perfectly good fallback, so never fail the run over a screenshot.
        return None


def _format_tables(tables: Iterable[Iterable[Iterable[object]]]) -> str:
    lines: list[str] = []
    for number, table in enumerate(tables, start=1):
        rows = [
            " | ".join("" if cell is None else str(cell).strip() for cell in row)
            for row in table
        ]
        rows = [row for row in rows if row.replace("|", "").strip()]
        if rows:
            lines.append(f"[Table {number}]\n" + "\n".join(rows))
    return "\n".join(lines)


# --------------------------------------------------------------------------- #
# PPTX
# --------------------------------------------------------------------------- #


def extract_pptx(path: str | Path, include_images: bool = True) -> DeckContent:
    """Extract text from every shape and embedded images from a PPTX."""
    from pptx import Presentation

    presentation = Presentation(str(path))
    slides = list(presentation.slides)
    slide_count = len(slides)
    wanted = _sample_indices(slide_count, MAX_IMAGES) if include_images else set()

    chunks: list[str] = []
    images: list[bytes] = []

    for index, slide in enumerate(slides):
        texts: list[str] = []
        blobs: list[bytes] = []
        for shape in _walk_shapes(slide.shapes):
            texts.extend(_shape_text(shape))
            if index in wanted:
                blob = _shape_image(shape)
                if blob is not None:
                    blobs.append(blob)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                texts.append(f"[Speaker notes] {notes.text.strip()}")

        chunks.append(_slide_block(index + 1, "\n".join(texts), ""))

        # One image per sampled slide keeps the cap meaningful on image-heavy decks.
        for blob in blobs[:1]:
            jpeg = _to_jpeg_bytes(blob)
            if jpeg is not None:
                images.append(jpeg)

    return DeckContent(
        text="\n\n".join(chunks), slide_count=slide_count, images=images[:MAX_IMAGES]
    )


def _walk_shapes(shapes) -> Iterable[object]:
    """Yield every shape, descending into groups."""
    for shape in shapes:
        if getattr(shape, "shape_type", None) is not None and hasattr(shape, "shapes"):
            yield from _walk_shapes(shape.shapes)
        else:
            yield shape


def _shape_text(shape) -> list[str]:
    out: list[str] = []
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()
        if text:
            out.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    return out


def _shape_image(shape) -> bytes | None:
    try:
        return shape.image.blob
    except (AttributeError, ValueError):
        return None


# --------------------------------------------------------------------------- #
# Shared helpers
# --------------------------------------------------------------------------- #


def _slide_block(number: int, body: str, extra: str) -> str:
    parts = [f"[Slide {number}]"]
    if body.strip():
        parts.append(body.strip())
    if extra.strip():
        parts.append(extra.strip())
    return "\n".join(parts)


def _sample_indices(total: int, limit: int) -> set[int]:
    """First + last + evenly sampled slide indices, capped at `limit`."""
    if total <= 0 or limit <= 0:
        return set()
    if total <= limit:
        return set(range(total))

    picked = {0, total - 1}
    remaining = limit - len(picked)
    if remaining > 0:
        step = (total - 1) / (remaining + 1)
        for i in range(1, remaining + 1):
            picked.add(min(total - 1, max(0, round(i * step))))
    # Rounding collisions can leave us short; backfill deterministically.
    for index in range(total):
        if len(picked) >= limit:
            break
        picked.add(index)
    return set(sorted(picked)[:limit])


def _to_jpeg_bytes(blob: bytes) -> bytes | None:
    from PIL import Image

    try:
        with Image.open(io.BytesIO(blob)) as image:
            return _to_jpeg(image)
    except Exception:
        return None


def _to_jpeg(image) -> bytes:
    """Downscale to MAX_IMAGE_WIDTH and encode as JPEG."""
    from PIL import Image

    working = image.convert("RGB")
    if working.width > MAX_IMAGE_WIDTH:
        height = max(1, round(working.height * MAX_IMAGE_WIDTH / working.width))
        working = working.resize((MAX_IMAGE_WIDTH, height), Image.LANCZOS)

    buffer = io.BytesIO()
    working.save(buffer, format="JPEG", quality=80, optimize=True)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Company name
# --------------------------------------------------------------------------- #

_NOISE_LINE = re.compile(
    r"^(confidential|private|proprietary|draft|do not distribute|"
    r"investor (deck|presentation)|pitch deck|seed|series [a-z]|"
    r"\d{4}|q[1-4]\s*\d{2,4})\b",
    re.IGNORECASE,
)


def guess_company_name(text: str, fallback: str = "Pitch Deck") -> str:
    """Best-effort company name from the deck's title slide.

    Used only for the report header — the analysis schema itself carries no
    company field, so this stays a presentation-layer heuristic.
    """
    first_slide = text.split("[Slide 2]", 1)[0]
    for raw in first_slide.splitlines():
        line = raw.strip().strip("|").strip()
        if not line or line.startswith("[Slide") or line.startswith("[Table"):
            continue
        if len(line) > 60 or _NOISE_LINE.match(line):
            continue
        if not any(char.isalpha() for char in line):
            continue
        return line
    return fallback
